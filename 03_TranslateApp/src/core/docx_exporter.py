"""Advanced DOCX exporter with format preservation and bilingual support."""

import logging
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)


class DocxExporter:
    """Handles exporting translations back to DOCX files."""

    def __init__(self, default_font: str = "Times New Roman"):
        self.default_font = default_font

    def export(
        self,
        original_path: str,
        output_path: str,
        pages_data: List[Dict[str, Any]],
        format_metadata: List[Dict[str, Any]],
        only_translated: bool = False,
        bilingual: bool = False
    ) -> None:
        """
        Export session data to a DOCX file.
        
        Args:
            original_path: Path to the template DOCX.
            output_path: Path to save the result.
            pages_data: List of page objects from session (containing 'jp' and 'vn' text).
            format_metadata: The list of metadata extracted by DocProcessor.
            only_translated: If True, only segments with non-empty 'vn' text are exported.
            bilingual: If True, exports both JP and VN text stacked.
        """
        doc = Document(original_path)
        
        # We need to map the flat list of translated segments back to the document structure.
        # The DocProcessor.extract_docx returns a flat list of strings.
        # Our session stores pages, which are groups of these strings joined by newlines.
        
        # 1. Flatten all translated texts from session pages
        all_vn_texts = []
        all_jp_texts = []
        for page in pages_data:
            # We assume pages were created by split('\n') of segments
            segments_vn = page["vn"].split("\n")
            segments_jp = page["jp"].split("\n")
            all_vn_texts.extend(segments_vn)
            all_jp_texts.extend(segments_jp)

        # 2. Iterate through metadata and apply changes to the doc
        # Important: Since we might be adding paragraphs (bilingual), 
        # we should be careful with indexing if we iterate and modify.
        # However, for standard replacement, it's straightforward.
        
        if not bilingual:
            self._export_standard(doc, all_vn_texts, format_metadata, only_translated)
        else:
            self._export_bilingual(doc, all_jp_texts, all_vn_texts, format_metadata, only_translated)

        doc.save(output_path)

    def _export_standard(self, doc: Document, vn_texts: List[str], metadata: List[Dict[str, Any]], only_translated: bool):
        """Replace JP text with VN text while keeping format."""
        for i, meta in enumerate(metadata):
            vn_text = vn_texts[i] if i < len(vn_texts) else ""
            
            if only_translated and not vn_text.strip():
                # If we only want translated, we might want to hide/remove untranslated?
                # For now, per spec clarification, we leave it empty.
                vn_text = ""

            para = self._get_paragraph(doc, meta)
            if para:
                self._apply_text_and_format(para, vn_text, meta["format"])

    def _export_bilingual(self, doc: Document, jp_texts: List[str], vn_texts: List[str], metadata: List[Dict[str, Any]], only_translated: bool):
        """Stack VN text below JP text."""
        # Note: Modifying document while iterating is tricky. 
        # For tables it's easier, but for main paragraphs we need to account for shifts.
        # Actually, adding a run to an existing paragraph is safer than adding a new paragraph.
        
        for i, meta in enumerate(metadata):
            jp_text = jp_texts[i] if i < len(jp_texts) else ""
            vn_text = vn_texts[i] if i < len(vn_texts) else ""
            
            if only_translated and not vn_text.strip():
                continue

            para = self._get_paragraph(doc, meta)
            if para:
                # Clear and add JP
                self._apply_text_and_format(para, jp_text, meta["format"])
                # Add newline and VN
                para.add_run("\n")
                self._apply_text_and_format(para, vn_text, meta["format"], is_translation=True)

    def _get_paragraph(self, doc: Document, meta: Dict[str, Any]) -> Optional[Paragraph]:
        """Utility to find a paragraph object based on metadata."""
        try:
            if meta["type"] == "paragraph":
                return doc.paragraphs[meta["index"]]
            elif meta["type"] == "table_cell":
                table = doc.tables[meta["table_index"]]
                cell = table.rows[meta["row_index"]].cells[meta["cell_index"]]
                return cell.paragraphs[meta["paragraph_index"]]
        except IndexError:
            logger.error(f"Index error while accessing document at {meta}")
        return None

    def _apply_text_and_format(self, para: Paragraph, text: str, fmt: Dict[str, Any], is_translation: bool = False):
        """Clear runs (if not translation) and add new run with styling."""
        if not is_translation:
            para.clear()
        
        run = para.add_run(text)
        
        if fmt.get("bold") is not None:
            run.bold = fmt["bold"]
        if fmt.get("italic") is not None:
            run.italic = fmt["italic"]
        if fmt.get("font_size"):
            run.font.size = Pt(fmt["font_size"])
        
        # Apply Vietnamese-friendly font
        run.font.name = self.default_font
        
        if fmt.get("font_color"):
            try:
                run.font.color.rgb = RGBColor.from_string(fmt["font_color"])
            except:
                pass
