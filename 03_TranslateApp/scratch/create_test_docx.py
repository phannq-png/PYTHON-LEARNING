from docx import Document

def create_test_docx(path):
    doc = Document()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('【Introduction】')
    doc.add_paragraph('This is the first segment of the medical report.')
    doc.add_paragraph('システム means system.')
    doc.add_paragraph('【Methods】')
    doc.add_paragraph('We used a diagnostic tool.')
    doc.add_paragraph('診断 is diagnosis.')
    doc.save(path)
    print(f"Created {path}")

if __name__ == "__main__":
    create_test_docx("test_input.docx")
