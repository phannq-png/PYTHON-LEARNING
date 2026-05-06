import os
import unittest
from docx import Document
from docx.shared import Pt, RGBColor

from src.core.doc_processor import DocumentProcessor

class TestDocumentProcessor(unittest.TestCase):
    """Test suite for DocumentProcessor class."""

    def setUp(self) -> None:
        """Setup test environment by creating a dummy DOCX file."""
        self.processor = DocumentProcessor()
        self.test_dir = "tests/test_data"
        if not os.path.exists(self.test_dir):
            os.makedirs(self.test_dir)
            
        self.input_file = os.path.join(self.test_dir, "sample.docx")
        self.output_file = os.path.join(self.test_dir, "output.docx")
        
        # Create a sample document
        doc = Document()
        
        # Add normal paragraph
        doc.add_paragraph("This is a normal paragraph.")
        
        # Add paragraph with specific format
        p2 = doc.add_paragraph()
        run = p2.add_run("This is bold and red.")
        run.bold = True
        run.font.color.rgb = RGBColor(255, 0, 0)
        run.font.size = Pt(14)
        
        # Add a table
        table = doc.add_table(rows=1, cols=2)
        cell1 = table.cell(0, 0)
        cell1.text = "Cell 1 text"
        cell2 = table.cell(0, 1)
        # Apply italic to cell 2
        cell2.paragraphs[0].clear()
        r2 = cell2.paragraphs[0].add_run("Cell 2 italic")
        r2.italic = True
        
        doc.save(self.input_file)

    def tearDown(self) -> None:
        """Clean up test files."""
        if os.path.exists(self.input_file):
            os.remove(self.input_file)
        if os.path.exists(self.output_file):
            os.remove(self.output_file)

    def test_extract_docx(self) -> None:
        """Test extraction of text and metadata."""
        texts, metadata = self.processor.extract_docx(self.input_file)
        
        self.assertEqual(len(texts), 4)
        self.assertEqual(texts[0], "This is a normal paragraph.")
        self.assertEqual(texts[1], "This is bold and red.")
        self.assertEqual(texts[2], "Cell 1 text")
        self.assertEqual(texts[3], "Cell 2 italic")
        
        # Check metadata for normal paragraph
        self.assertFalse(metadata[0]["format"]["bold"])
        
        # Check metadata for formatted paragraph
        self.assertTrue(metadata[1]["format"]["bold"])
        self.assertEqual(metadata[1]["format"]["font_color"], "FF0000")
        self.assertEqual(metadata[1]["format"]["font_size"], 14.0)
        
        # Check metadata for italic cell
        self.assertTrue(metadata[3]["format"]["italic"])

    def test_export_docx(self) -> None:
        """Test exporting translated text with original format."""
        texts, metadata = self.processor.extract_docx(self.input_file)
        
        # Mock translated texts
        translated = [
            "Đây là đoạn văn bình thường.",
            "Đoạn này in đậm và màu đỏ.",
            "Nội dung ô 1",
            "Nội dung ô 2 in nghiêng"
        ]
        
        self.processor.export_docx(self.input_file, translated, metadata, self.output_file)
        
        self.assertTrue(os.path.exists(self.output_file))
        
        # Verify the output document
        out_doc = Document(self.output_file)
        
        # Check first paragraph
        p1 = out_doc.paragraphs[0]
        self.assertEqual(p1.text, "Đây là đoạn văn bình thường.")
        if p1.runs:
            self.assertEqual(p1.runs[0].font.name, "Times New Roman")
            
        # Check second paragraph formatting
        p2 = out_doc.paragraphs[1]
        self.assertEqual(p2.text, "Đoạn này in đậm và màu đỏ.")
        self.assertTrue(p2.runs[0].bold)
        self.assertEqual(str(p2.runs[0].font.color.rgb), "FF0000")
        self.assertEqual(p2.runs[0].font.size.pt, 14.0)
        self.assertEqual(p2.runs[0].font.name, "Times New Roman")
        
        # Check table
        table = out_doc.tables[0]
        cell2_para = table.cell(0, 1).paragraphs[0]
        self.assertEqual(cell2_para.text, "Nội dung ô 2 in nghiêng")
        self.assertTrue(cell2_para.runs[0].italic)
        self.assertEqual(cell2_para.runs[0].font.name, "Times New Roman")

    def test_extract_invalid_file(self) -> None:
        """Test extracting from a non-existent file."""
        with self.assertRaises(FileNotFoundError):
            self.processor.extract_docx("non_existent.docx")

if __name__ == '__main__':
    unittest.main()
